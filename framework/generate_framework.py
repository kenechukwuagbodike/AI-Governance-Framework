"""
Renders framework_content.py into the AI Governance Framework .docx.

Run from the framework/ folder with the project venv active:
    python generate_framework.py

Output: ai_governance_framework.docx, gitignored, in this same folder.
"""

import logging
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

import framework_content as content

logging.basicConfig(level=logging.INFO, format="%(message)s")
logger = logging.getLogger(__name__)

NAVY = RGBColor(0x14, 0x21, 0x3D)
GREY = RGBColor(0x66, 0x66, 0x66)
DARK = RGBColor(0x1A, 0x1A, 0x1A)

HEADING_FONT = "Cambria"
BODY_FONT = "Georgia"

OUTPUT_PATH = Path(__file__).parent / "ai_governance_framework.docx"


def set_base_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    style_names = [s.name for s in document.styles]
    for list_style_name in ("List Bullet", "List Number"):
        if list_style_name in style_names:
            list_style = document.styles[list_style_name]
            list_style.font.name = BODY_FONT
            list_style.font.size = Pt(11)
            list_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    heading1 = document.styles["Heading 1"]
    heading1.font.name = HEADING_FONT
    heading1.font.size = Pt(20)
    heading1.font.bold = True
    heading1.font.color.rgb = NAVY
    heading1.paragraph_format.space_before = Pt(0)
    heading1.paragraph_format.space_after = Pt(4)

    heading2 = document.styles["Heading 2"]
    heading2.font.name = HEADING_FONT
    heading2.font.size = Pt(13)
    heading2.font.bold = True
    heading2.font.color.rgb = NAVY
    heading2.paragraph_format.space_before = Pt(14)
    heading2.paragraph_format.space_after = Pt(4)


def add_title_page(document: Document) -> None:
    document.add_paragraph()
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(content.FRAMEWORK_TITLE)
    run.font.name = HEADING_FONT
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = NAVY

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(content.FRAMEWORK_SUBTITLE)
    run.font.name = BODY_FONT
    run.font.size = Pt(13)
    run.font.color.rgb = GREY

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Prepared by Kene Agbodike  |  {date.today():%B %Y}")
    run.font.size = Pt(10)
    run.font.color.rgb = GREY

    document.add_page_break()


def add_big_idea(document: Document, big_idea: str, label: str = "Big idea") -> None:
    para = document.add_paragraph()
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(12)
    para.paragraph_format.left_indent = Cm(0.4)
    label_run = para.add_run(f"{label}: ")
    label_run.font.name = HEADING_FONT
    label_run.font.bold = True
    label_run.font.color.rgb = NAVY
    label_run.font.size = Pt(11)
    body_run = para.add_run(big_idea)
    body_run.font.name = BODY_FONT
    body_run.font.italic = True
    body_run.font.size = Pt(11)
    body_run.font.color.rgb = DARK


def add_source_line(document: Document, backbone: str, supporting_refs: str,
                     nist_function: str) -> None:
    para = document.add_paragraph()
    para.paragraph_format.space_after = Pt(10)
    label_run = para.add_run(f"NIST function: {nist_function}   |   Backbone: ")
    label_run.font.size = Pt(9)
    label_run.font.color.rgb = GREY
    backbone_run = para.add_run(backbone)
    backbone_run.font.size = Pt(9)
    backbone_run.font.bold = True
    backbone_run.font.color.rgb = GREY
    supporting_run = para.add_run(f"   |   Supporting: {supporting_refs}")
    supporting_run.font.size = Pt(9)
    supporting_run.font.color.rgb = GREY


def add_introduction(document: Document) -> None:
    document.add_heading("Introduction", level=1)
    add_big_idea(document, content.INTRODUCTION["big_idea"])
    for paragraph_text in content.INTRODUCTION["body"]:
        document.add_paragraph(paragraph_text)
    document.add_page_break()


def add_subsection(document: Document, subsection: dict) -> None:
    document.add_heading(subsection["heading"], level=2)
    for paragraph_text in subsection.get("body", []):
        document.add_paragraph(paragraph_text)
    for bullet_text in subsection.get("bullets", []):
        document.add_paragraph(bullet_text, style="List Bullet")


def add_section(document: Document, section: dict) -> None:
    document.add_heading(f"Section {section['number']}: {section['title']}", level=1)
    add_source_line(
        document,
        section["backbone"],
        section["supporting_refs"],
        section["nist_function"],
    )
    add_big_idea(document, section["big_idea"])
    for subsection in section["subsections"]:
        add_subsection(document, subsection)
    document.add_page_break()


def add_references(document: Document) -> None:
    document.add_heading("References", level=1)
    for reference in content.REFERENCES:
        document.add_paragraph(reference, style="List Number")


def build_document() -> Document:
    document = Document()
    set_base_styles(document)
    add_title_page(document)
    add_introduction(document)
    for section in content.FRAMEWORK_SECTIONS:
        logger.info("Rendering Section %s: %s", section["number"], section["title"])
        add_section(document, section)
    add_references(document)
    return document


def main() -> None:
    document = build_document()
    document.save(OUTPUT_PATH)
    logger.info("Saved %s", OUTPUT_PATH)


if __name__ == "__main__":
    main()
